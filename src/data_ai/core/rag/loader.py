"""
文档加载器

支持从多种格式加载文档内容。
"""

from __future__ import annotations

import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ...utils.logging import get_logger

logger = get_logger(__name__)


@dataclass
class Document:
    """文档"""

    content: str
    metadata: dict[str, Any] = field(default_factory=dict)
    page: int | None = None


class BaseLoader(ABC):
    """加载器基类"""

    @abstractmethod
    def load(self, source: str) -> list[Document]:
        """加载文档"""
        pass

    def load_many(self, sources: list[str]) -> list[Document]:
        """批量加载"""
        docs = []
        for source in sources:
            docs.extend(self.load(source))
        return docs


class TextLoader(BaseLoader):
    """纯文本加载器

    从文本文件加载文档内容。
    """

    def __init__(self, encoding: str = "utf-8"):
        self.encoding = encoding

    def load(self, source: str) -> list[Document]:
        """加载文本文件

        Args:
            source: 文件路径

        Returns:
            文档列表
        """
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {source}")

        with open(path, encoding=self.encoding) as f:
            content = f.read()

        return [
            Document(
                content=content,
                metadata={"source": str(path.absolute()), "format": "text", "type": "text"},
            )
        ]


class PDFLoader(BaseLoader):
    """PDF 加载器

    从PDF文件加载文档内容，支持pypdf和PyMuPDF后端。
    """

    def __init__(self, extract_images: bool = False):
        self.extract_images = extract_images

    def load(self, source: str) -> list[Document]:
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {source}")

        docs = []

        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))

            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(
                        Document(
                            content=text.strip(),
                            metadata={
                                "source": str(path.absolute()),
                                "format": "pdf",
                                "page": i + 1,
                                "total_pages": len(reader.pages),
                            },
                            page=i + 1,
                        )
                    )
        except ImportError:
            try:
                import fitz

                doc = fitz.open(str(path))
                for i, page in enumerate(doc):
                    text = page.get_text()
                    if text.strip():
                        docs.append(
                            Document(
                                content=text.strip(),
                                metadata={
                                    "source": str(path.absolute()),
                                    "format": "pdf",
                                    "page": i + 1,
                                    "total_pages": len(doc),
                                },
                                page=i + 1,
                            )
                        )
                doc.close()
            except ImportError:
                raise ImportError("PDF 加载需要 pypdf 或 PyMuPDF: pip install pypdf 或 pip install pymupdf")

        return docs


class MarkdownLoader(BaseLoader):
    """Markdown 加载器"""

    def __init__(self, split_by_heading: bool = False):
        """初始化Markdown加载器

        Args:
            split_by_heading: 是否按标题分割文档，默认为False返回单个文档
        """
        self.split_by_heading = split_by_heading

    def load(self, source: str) -> list[Document]:
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {source}")

        with open(path, encoding="utf-8") as f:
            content = f.read()

        # 提取所有标题
        headers = []
        for match in re.finditer(r"^(#{1,6})\s+(.+)$", content, re.MULTILINE):
            headers.append(match.group(2).strip())

        if self.split_by_heading:
            # 按标题分割为多个文档段
            docs = []
            sections = re.split(r"^(#{1,6}\s+.+)$", content, flags=re.MULTILINE)

            current_section = ""
            current_heading = ""

            for section in sections:
                section = section.strip()
                if not section:
                    continue

                if re.match(r"^#{1,6}\s+", section):
                    if current_section.strip():
                        docs.append(
                            Document(
                                content=current_section.strip(),
                                metadata={
                                    "source": str(path.absolute()),
                                    "format": "markdown",
                                    "type": "markdown",
                                    "heading": current_heading,
                                    "headers": headers,
                                },
                            )
                        )
                    current_heading = section
                    current_section = section + "\n"
                else:
                    current_section += section + "\n"

            if current_section.strip():
                docs.append(
                    Document(
                        content=current_section.strip(),
                        metadata={
                            "source": str(path.absolute()),
                            "format": "markdown",
                            "type": "markdown",
                            "heading": current_heading,
                            "headers": headers,
                        },
                    )
                )

            return (
                docs
                if docs
                else [
                    Document(
                        content=content,
                        metadata={
                            "source": str(path.absolute()),
                            "format": "markdown",
                            "type": "markdown",
                            "headers": headers,
                        },
                    )
                ]
            )
        else:
            # 返回单个文档
            return [
                Document(
                    content=content,
                    metadata={
                        "source": str(path.absolute()),
                        "format": "markdown",
                        "type": "markdown",
                        "headers": headers,
                    },
                )
            ]


class HTMLLoader(BaseLoader):
    """HTML 加载器"""

    def __init__(self, extract_text: bool = True, strip_tags: bool = False):
        self.extract_text = extract_text
        self.strip_tags = strip_tags

    def load(self, source: str) -> list[Document]:
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {source}")

        with open(path, encoding="utf-8") as f:
            html = f.read()

        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(html, "html.parser")

            # 移除脚本和样式
            for tag in soup(["script", "style"]):
                tag.decompose()

            text = soup.get_text(separator="\n", strip=True)
            title = soup.title.string if soup.title else ""

        except ImportError:
            text = re.sub(r"<[^>]+>", " ", html)
            text = re.sub(r"\s+", " ", text).strip()
            title = ""

        return [
            Document(
                content=text,
                metadata={
                    "source": str(path.absolute()),
                    "format": "html",
                    "type": "html",
                    "title": title,
                },
            )
        ]


class DocxLoader(BaseLoader):
    """Word 文档加载器"""

    def load(self, source: str) -> list[Document]:
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {source}")

        try:
            from docx import Document

            doc = Document(str(path))

            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n\n".join(paragraphs)

            return [
                Document(
                    content=content,
                    metadata={
                        "source": str(path.absolute()),
                        "format": "docx",
                        "title": doc.core_properties.title or "",
                        "author": doc.core_properties.author or "",
                    },
                )
            ]
        except ImportError:
            raise ImportError("Word 加载需要 python-docx: pip install python-docx")


class DocumentLoader:
    """智能文档加载器，根据文件扩展名自动选择加载器"""

    LOADERS = {
        ".txt": TextLoader,
        ".md": MarkdownLoader,
        ".markdown": MarkdownLoader,
        ".pdf": PDFLoader,
        ".html": HTMLLoader,
        ".htm": HTMLLoader,
        ".docx": DocxLoader,
    }

    @classmethod
    def load(cls, source: str, **kwargs) -> list[Document]:
        """自动检测格式并加载"""
        path = Path(source)
        suffix = path.suffix.lower()

        loader_class = cls.LOADERS.get(suffix)
        if not loader_class:
            raise ValueError(f"不支持的文件格式: {suffix}")

        loader = loader_class(**kwargs)
        return loader.load(source)

    @classmethod
    def load_directory(cls, directory: str, recursive: bool = True, **kwargs) -> list[Document]:
        """加载目录下所有支持的文档"""
        dir_path = Path(directory)
        if not dir_path.exists():
            raise FileNotFoundError(f"目录不存在: {directory}")

        docs = []
        pattern = "**/*" if recursive else "*"

        for file_path in dir_path.glob(pattern):
            if file_path.is_file() and file_path.suffix.lower() in cls.LOADERS:
                try:
                    file_docs = cls.load(str(file_path), **kwargs)
                    docs.extend(file_docs)
                    logger.info(f"已加载: {file_path.name} ({len(file_docs)} 个文档段)")
                except Exception as e:
                    logger.warning(f"加载失败: {file_path.name}: {e}")

        return docs
